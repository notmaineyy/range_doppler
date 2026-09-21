"""
Generate clean result figures and build the Task 1 slide deck.

Run::

    .venv/bin/python make_task1_slides.py

Outputs:
    outputs/slides/*.png        result figures
    task1_slides.pptx           presentation
"""

import os

import numpy as np
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, Polygon, FancyBboxPatch
from matplotlib.patches import Arc

import radar
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "outputs", "slides")
os.makedirs(FIG, exist_ok=True)

DPI = 150
DARK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x55, 0x55, 0x55)


# --------------------------------------------------------------------------- #
# Simulation helper
# --------------------------------------------------------------------------- #
def compute(B=200e6, vr=0.0, La=2.0, vt=0.0):
    cfg = dict(radar.make_cfg())
    cfg["B"] = B
    cfg["La"] = La
    d = radar.derive(cfg)
    t_m = radar.slow_time(cfg, d)
    t_f, _ = radar.fast_time(cfg)
    R = radar.slant_range(cfg, t_m, vr=vr, vt=vt)
    data = radar.raw_echo(cfg, d, t_m, t_f, R)
    rc = radar.range_compress(cfg, d, data)
    rd_map, ac = radar.rda_compress(cfg, d, rc)
    return cfg, d, t_m, t_f, rd_map, ac


# --------------------------------------------------------------------------- #
# Figures
# --------------------------------------------------------------------------- #
def fig_geometry():
    fig, ax = plt.subplots(figsize=(8, 4.6))
    ax.set_xlim(0, 10)
    ax.set_ylim(-0.6, 6.2)
    ax.axis("off")

    # ground / swath
    ax.plot([0, 10], [0, 0], color="#8a6d3b", lw=6, solid_capstyle="butt")
    ax.text(0.15, 0.18, "ground", ha="left", fontsize=9, color="#8a6d3b")

    # platform
    ax.add_patch(FancyBboxPatch((4.2, 5.0), 1.6, 0.5, boxstyle="round,pad=0.05",
                                fc="#2c3e50", ec="none"))
    ax.text(5.0, 5.25, "SAR platform", ha="center", va="center",
            color="white", fontsize=10, fontweight="bold")
    # velocity arrow
    ax.add_patch(FancyArrowPatch((6.0, 5.25), (7.6, 5.25),
                                 arrowstyle="-|>", mutation_scale=18, lw=2.5,
                                 color="#c0392b"))
    ax.text(6.8, 5.5, r"$v_p$  (along-track)", color="#c0392b", fontsize=11)

    # beam
    ax.add_patch(Polygon([[4.6, 5.0], [7.4, 5.0], [8.6, 0.0], [3.4, 0.0]],
                         closed=True, fc="#3498db", alpha=0.15, ec="#3498db", lw=1.5))
    ax.add_patch(Arc((5.0, 5.0), 1.6, 1.0, theta1=270, theta2=315,
                     color="#3498db", lw=1.5))
    ax.text(5.75, 4.55, r"beam $\theta=\lambda/L_a$", color="#2c7fb8", fontsize=9)

    # slant range arrow
    ax.add_patch(FancyArrowPatch((5.0, 5.0), (6.4, 0.0),
                                 arrowstyle="-|>", mutation_scale=14, lw=2,
                                 color="#27ae60", linestyle="--"))
    ax.text(5.2, 2.6, r"slant range $R_0$", color="#1e8449", fontsize=11, rotation=-72)

    # target
    ax.plot([6.4], [0.0], marker="o", ms=9, color="#c0392b")
    ax.text(6.4, 0.25, "point target", ha="center", color="#c0392b", fontsize=10)

    # axis labels
    ax.annotate("", xy=(9.6, -0.05), xytext=(7.0, -0.05),
                arrowprops=dict(arrowstyle="<->", color="#555"))
    ax.text(8.3, -0.42, "azimuth / along-track", ha="center", fontsize=9, color="#555")

    ax.set_title("SAR geometry: a moving antenna illuminates a target", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "geometry.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_pipeline():
    fig, ax = plt.subplots(figsize=(13, 2.6))
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 2.4)
    ax.axis("off")

    steps = [
        "Point\ntarget",
        "Raw echo\n(slow x fast time)",
        "Range\ncompression",
        "Azimuth FFT\n= Range-Doppler map",
        "Azimuth\nmatched filter\n(RDA)",
        "Compressed\nSAR image",
    ]
    x = 0.2
    w = 1.85
    for i, s in enumerate(steps):
        ax.add_patch(FancyBboxPatch((x, 0.7), w, 1.0, boxstyle="round,pad=0.06",
                                    fc="#eaf2fb", ec="#2c7fb8", lw=1.5))
        ax.text(x + w / 2, 1.2, s, ha="center", va="center", fontsize=10)
        if i < len(steps) - 1:
            ax.add_patch(FancyArrowPatch((x + w, 1.2), (x + w + 0.28, 1.2),
                                         arrowstyle="-|>", mutation_scale=16,
                                         lw=2, color="#555"))
        x += w + 0.28

    ax.text(0.2, 0.15, "each stage is a step in the Range-Doppler Algorithm (RDA)",
            fontsize=9, color="#555", style="italic")
    fig.tight_layout()
    p = os.path.join(FIG, "pipeline.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_baseline():
    cfg, d, t_m, t_f, rd_map, ac = compute()
    f_a, R_axis, x_axis, _, Na, Nr = radar.axes(cfg, d)

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.6))
    im0 = axs[0].imshow(radar.to_dB(rd_map), aspect="auto", origin="lower",
                        cmap="inferno", vmin=-60, vmax=0,
                        extent=[R_axis[0], R_axis[-1], f_a[0], f_a[-1]])
    axs[0].axhline(0, color="w", ls=":", lw=0.8)
    axs[0].set_title("Range-Doppler map")
    axs[0].set_xlabel("Range [m]")
    axs[0].set_ylabel("Doppler frequency [Hz]")
    fig.colorbar(im0, ax=axs[0], fraction=0.046, pad=0.04, label="dB")

    im1 = axs[1].imshow(radar.to_dB(ac), aspect="auto", origin="lower",
                        cmap="inferno", vmin=-60, vmax=0,
                        extent=[R_axis[0], R_axis[-1], x_axis[0], x_axis[-1]])
    axs[1].set_title("Compressed image (range x azimuth)")
    axs[1].set_xlabel("Range [m]")
    axs[1].set_ylabel("Azimuth $x=v_p t$ [m]")
    fig.colorbar(im1, ax=axs[1], fraction=0.046, pad=0.04, label="dB")

    fig.suptitle(f"Baseline: dR={d['dR']:.2f} m, rho_az={d['rho_az']:.1f} m, "
                 f"Bd={d['Bd']:.0f} Hz", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "baseline.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def _range_slice(ac, R_axis):
    j, _ = np.unravel_index(np.argmax(np.abs(ac)), ac.shape)
    s = np.abs(ac[j, :])
    return s / s.max()


def _azimuth_slice(ac, x_axis):
    _, i = np.unravel_index(np.argmax(np.abs(ac)), ac.shape)
    s = np.abs(ac[:, i])
    return s / s.max()


def fig_bandwidth():
    Bs = [50e6, 200e6, 500e6]
    fig, axs = plt.subplots(1, 3, figsize=(13, 4))
    for ax, B in zip(axs, Bs):
        cfg, d, t_m, t_f, rd_map, ac = compute(B=B)
        _, R_axis, _, _, _, _ = radar.axes(cfg, d)
        s = _range_slice(ac, R_axis)
        ax.plot(R_axis, 20 * np.log10(s + 1e-9), lw=1.3)
        ax.axvline(5000, color="r", ls="--", lw=0.9)
        ax.set_title(f"B = {B/1e6:.0f} MHz   (dR = {radar.C/(2*B):.2f} m)")
        ax.set_xlabel("Range [m]")
        ax.set_ylabel("Normalised power [dB]")
        ax.set_xlim(4985, 5015)
        ax.set_ylim(-40, 2)
        ax.grid(alpha=0.3)
    fig.suptitle("Bandwidth B sets RANGE clarity (dR = c/2B); target stays at R0",
                 fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "bandwidth.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_radial():
    vrs = [0.0, 1.0, 2.0]
    fig, axs = plt.subplots(1, 3, figsize=(14, 4.3))
    for ax, vr in zip(axs, vrs):
        cfg, d, t_m, t_f, rd_map, ac = compute(vr=vr)
        f_a, R_axis, _, _, _, _ = radar.axes(cfg, d)
        fd = 2.0 * vr / d["lam"]
        ax.imshow(radar.to_dB(rd_map), aspect="auto", origin="lower",
                  cmap="inferno", vmin=-60, vmax=0,
                  extent=[R_axis[0], R_axis[-1], f_a[0], f_a[-1]])
        ax.axhline(fd, color="cyan", ls="--", lw=1.0)
        ax.axhline(0, color="w", ls=":", lw=0.7)
        ax.set_title(f"$v_r$ = {vr:.0f} m/s   ($f_{{dc}}$ = +{fd:.0f} Hz)")
        ax.set_xlabel("Range [m]")
        ax.set_ylabel("Doppler [Hz]")
    fig.suptitle("Radial velocity shifts the Doppler centroid (f_dc = 2vr/lambda) "
                 "-> position error", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "radial.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_antenna():
    Las = [1.0, 2.0, 4.0]
    fig, axs = plt.subplots(1, 3, figsize=(13, 4))
    for ax, La in zip(axs, Las):
        cfg, d, t_m, t_f, rd_map, ac = compute(La=La)
        _, _, x_axis, _, _, _ = radar.axes(cfg, d)
        s = _azimuth_slice(ac, x_axis)
        ax.plot(x_axis, 20 * np.log10(s + 1e-9), lw=1.3)
        ax.axvline(0, color="r", ls="--", lw=0.9)
        ax.set_title(f"$L_a$ = {La:.0f} m   (rho_az = {La/2:.1f} m, "
                     f"Bd = {2*150/La:.0f} Hz)")
        ax.set_xlabel("Azimuth [m]")
        ax.set_ylabel("Normalised power [dB]")
        ax.set_xlim(-25, 25)
        ax.set_ylim(-40, 2)
        ax.grid(alpha=0.3)
    fig.suptitle("Antenna length La sets AZIMUTH clarity (rho_az = La/2); "
                 "target stays at x0", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "antenna.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_alongtrack():
    vts = [0.0, 40.0, 80.0]
    fig, axs = plt.subplots(1, 3, figsize=(13, 4))
    for ax, vt in zip(axs, vts):
        cfg, d, t_m, t_f, rd_map, ac = compute(vt=vt)
        _, _, x_axis, _, _, _ = radar.axes(cfg, d)
        s = _azimuth_slice(ac, x_axis)
        ax.plot(x_axis, 20 * np.log10(s + 1e-9), lw=1.3)
        ax.axvline(0, color="r", ls="--", lw=0.9)
        ax.set_title(f"target $v_t$ = {vt:.0f} m/s   (filter mismatch)")
        ax.set_xlabel("Azimuth [m]")
        ax.set_ylabel("Normalised power [dB]")
        ax.set_xlim(-80, 80)
        ax.set_ylim(-40, 2)
        ax.grid(alpha=0.3)
    fig.suptitle("Target along-track velocity mismatches the azimuth filter "
                 "-> defocus (clarity loss)", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "alongtrack.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_snr():
    cfg = radar.make_cfg()
    t_f, _ = radar.fast_time(cfg)
    cases = [("clean", 40.0, False), ("SNR = 25 dB", 25.0, False),
             ("SNR = 10 dB", 10.0, False), ("speckle (SNR 25 dB)", 25.0, True)]
    fig, axs = plt.subplots(2, 2, figsize=(11, 8))
    for ax, (title, snr, speckle) in zip(axs.ravel(), cases):
        rng = np.random.default_rng(0)
        _, im_disp, meta = radar.simulate(cfg, t_f, 200e6, 0.0, 2.0, 0.0,
                                          snr, speckle, rng)
        im = ax.imshow(im_disp, aspect="auto", origin="lower", cmap="inferno",
                       vmin=-45, vmax=0,
                       extent=[meta["R_axis"][0], meta["R_axis"][-1],
                               meta["x_axis"][0], meta["x_axis"][-1]])
        ax.plot([5000], [0], marker="o", ms=16, mfc="none", mec="cyan", mew=1.6)
        ax.set_title(title)
        ax.set_xlabel("Range [m]")
        ax.set_ylabel("Azimuth [m]")
        fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04, label="dB")
    fig.suptitle("Noise and speckle reduce clarity / detectability "
                 "(cyan circle = true target position)", fontsize=12)
    fig.tight_layout()
    p = os.path.join(FIG, "snr_speckle.png")
    fig.savefig(p, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return p


# --------------------------------------------------------------------------- #
# Slide deck
# --------------------------------------------------------------------------- #
def _bg(slide, prs):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK


def _body(slide, items, left=0.7, top=1.35, width=12.0, height=5.7, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                                  Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        text, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text
        p.level = level
        p.font.size = Pt(size - 2 * level)
        p.font.color.rgb = DARK if level == 0 else GREY
        p.space_after = Pt(6)
    return tb


def _picture(slide, path, top=1.35, width=11.6, caption=None):
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    left = (13.333 - width) / 2.0
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width))
    if caption:
        cap_top = top + width * h / w + 0.08
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(cap_top), Inches(12.0),
                                      Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = caption
        p.font.size = Pt(13)
        p.font.color.rgb = GREY
        p.alignment = PP_ALIGN.CENTER
    return pic


def add_title_slide(prs, title, subtitle, author):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, prs)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(2.2), Inches(13.333), Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = DARK
    p = tf.add_paragraph(); p.text = subtitle
    p.font.size = Pt(20); p.font.color.rgb = GREY
    p = tf.add_paragraph(); p.text = author
    p.font.size = Pt(14); p.font.color.rgb = GREY
    return slide


def add_bullets(prs, title, items, size=18):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, prs)
    _title(slide, title)
    _body(slide, items, size=size)
    return slide


def add_picture_slide(prs, title, path, bullets=None, caption=None, width=9.6):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, prs)
    _title(slide, title)
    if bullets:
        _body(slide, bullets, left=0.6, top=1.35, width=3.7, height=5.6, size=14)
        _picture(slide, path, top=1.5, width=8.6, caption=caption)
        # shift picture right
        for shp in slide.shapes:
            if shp.shape_type == 13:  # PICTURE
                shp.left = Inches(4.6)
    else:
        _picture(slide, path, top=1.5, width=width, caption=caption)
    return slide


def add_table(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, prs)
    _title(slide, title)
    nrows, ncols = len(rows) + 1, len(headers)
    tbl = slide.shapes.add_table(nrows, ncols, Inches(0.6), Inches(1.5),
                                 Inches(12.1), Inches(0.6 * nrows)).table
    for c, htext in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = htext
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(14); para.font.bold = True
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid(); cell.fill.fore_color.rgb = DARK
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(13)
            if r % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF2, 0xF5, 0xFA)
    return slide


def build_deck(figs):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "SAR Range-Doppler Explorer",
        "Range, azimuth and Doppler in a Synthetic Aperture Radar image",
        "Task 1 - interactive simulation (range_doppler_interactive.py)",
    )

    add_bullets(prs, "What is SAR?", [
        "SAR = Synthetic Aperture Radar: a radar on a moving platform images the ground.",
        ("A small antenna is moved along a path; the many echoes are combined to act "
         "like one very long antenna - a 'synthetic aperture'.", 1),
        "The radar sends short frequency-swept pulses (chirps) and listens for echoes.",
        "Echo delay gives distance (range); the changing echo frequency (Doppler) "
        "from motion gives the along-track position (azimuth).",
        "Result: a two-dimensional image that works day or night and through clouds.",
    ])

    add_picture_slide(
        prs, "Two dimensions of a SAR image", figs["geometry"],
        bullets=[
            "Range (cross-track):",
            ("measured from echo delay, R = c*tau/2", 1),
            ("clarity set by bandwidth B", 1),
            "Azimuth (along-track):",
            ("measured from Doppler history", 1),
            ("clarity set by antenna length La", 1),
        ],
        width=7.0,
    )

    add_bullets(prs, "What the interactive explorer does", [
        "Simulates a single point target and processes it with the Range-Doppler Algorithm.",
        "Left panel: the Range-Doppler map (range vs Doppler frequency).",
        "Right panel: the fully compressed SAR image (range vs azimuth).",
        "Five sliders + one checkbox update both panels live (~12 frames/second).",
        "Live read-out: dR, rho_az, f_dc, Bd, Na.",
        "Press 'r' or click Reset to return to defaults.",
    ])

    add_picture_slide(
        prs, "The processing chain (methodology)", figs["pipeline"],
        bullets=[
            "1. Build the raw echo in slow-time x fast-time.",
            "2. Range compression (matched filter).",
            "3. FFT along azimuth -> Range-Doppler map.",
            "4. Azimuth matched filter (in Doppler domain).",
            "5. Display in dB.",
        ],
        width=11.0,
    )

    add_picture_slide(
        prs, "Baseline result", figs["baseline"],
        bullets=[
            "A point target focuses to a bright spot.",
            "Range-Doppler map: energy in a narrow Doppler band at R0.",
            "Compressed image: one sharp point at (R0, x0).",
        ],
        width=11.6,
    )

    add_picture_slide(
        prs, "Variable 1 - Bandwidth B", figs["bandwidth"],
        bullets=[
            "Range resolution dR = c / (2B).",
            "Wider B -> narrower mainlobe -> sharper in range.",
            "Target POSITION does not move: peak stays at R0.",
            "Clarity only, not position.",
        ],
        width=9.6,
    )

    add_picture_slide(
        prs, "Variable 2 - Radial velocity vr", figs["radar"] if False else figs["radial"],
        bullets=[
            "Doppler centroid f_dc = 2 vr / lambda.",
            "Shifts the target along the Doppler axis.",
            "Also causes range walk (vr * Ta).",
            "Uncompensated -> azimuth position error and smearing.",
        ],
        width=9.6,
    )

    add_picture_slide(
        prs, "Variable 3 - Antenna length La", figs["antenna"],
        bullets=[
            "Azimuth resolution rho_az = La / 2.",
            "Doppler bandwidth Bd = 2 vp / La.",
            "Shorter antenna -> wider beam -> finer azimuth resolution.",
            "Target POSITION does not move: peak stays at x0.",
        ],
        width=9.6,
    )

    add_picture_slide(
        prs, "Variable 4 - Target along-track velocity vt", figs["alongtrack"],
        bullets=[
            "A moving target's Doppler rate differs from the stationary filter.",
            "Phase mismatch -> azimuth defocus (smeared, lower peak).",
            "Clarity loss; position roughly unchanged.",
        ],
        width=9.6,
    )

    add_picture_slide(
        prs, "Variable 5 - SNR and speckle", figs["snr_speckle"],
        bullets=[
            "Low SNR -> noise floor rises, target buried.",
            "Speckle = multiplicative Rayleigh noise (grainy texture).",
            "Both reduce detectability / clarity.",
            "Position unaffected.",
        ],
        width=9.6,
    )

    add_table(
        prs, "Results at a glance: what affects clarity vs position",
        ["Variable", "Affects clarity", "Affects position", "Key formula"],
        [
            ["Bandwidth B", "Range clarity", "No", "dR = c / 2B"],
            ["Radial velocity vr", "Defocus if large", "Doppler / azimuth", "f_dc = 2vr/lambda"],
            ["Antenna length La", "Azimuth clarity", "No", "rho_az = La/2"],
            ["Along-track velocity vt", "Azimuth defocus", "~No", "Doppler-rate mismatch"],
            ["SNR / speckle", "Detectability", "No", "noise floor / Rayleigh"],
        ],
    )

    add_bullets(prs, "Cheat sheet: the equations", [
        "Range resolution:            dR = c / (2B)",
        "Azimuth resolution:          rho_az = La / 2",
        "Doppler bandwidth:           Bd = 2 vp / La",
        "Doppler rate:                fR = 2 vp^2 / (lambda R0)",
        "Doppler centroid (moving):   f_dc = 2 vr / lambda",
        "Synthetic aperture time:     Ta = R0 * theta / vp,  theta = lambda / La",
        "Range to target:             R = c * tau / 2",
    ], size=19)

    add_bullets(prs, "Glossary of SAR terms", [
        "Range - distance to target, from echo delay.",
        "Azimuth / cross-range - along-track position, from Doppler.",
        "Slant range R0 - straight-line distance from radar to target.",
        "Chirp - a pulse whose frequency sweeps over time.",
        "Pulse / range compression - matched filtering that sharpens echoes in range.",
        "Doppler history - the changing echo frequency as the platform passes the target.",
        "Doppler centroid - the mean Doppler shift (zero for a still target broadside).",
        "Synthetic aperture - the long effective antenna formed by platform motion.",
        "PRF - pulses per second; must be high enough to sample the Doppler.",
        "RDA (Range-Doppler Algorithm) - process range first, then azimuth via Doppler.",
        "Speckle - grainy noise from many small scatterers in one resolution cell.",
    ], size=14)

    add_bullets(prs, "Key takeaways & how to run", [
        "Range clarity is set by BANDWIDTH; azimuth clarity by ANTENNA LENGTH.",
        "Radial velocity moves the target's Doppler/azimuth POSITION; large values defocus.",
        "Along-track velocity and noise degrade CLARITY, not position.",
        "The app applies the azimuth filter in the Doppler domain (the RDA), "
        "keeping updates real-time.",
        "Run it:  .venv/bin/python range_doppler_interactive.py",
        "Web version:  .venv/bin/streamlit run app.py",
    ], size=18)

    path = os.path.join(HERE, "task1_slides.pptx")
    prs.save(path)
    return path


def main():
    figs = {
        "geometry": fig_geometry(),
        "pipeline": fig_pipeline(),
        "baseline": fig_baseline(),
        "bandwidth": fig_bandwidth(),
        "radial": fig_radial(),
        "antenna": fig_antenna(),
        "alongtrack": fig_alongtrack(),
        "snr_speckle": fig_snr(),
    }
    for name, p in figs.items():
        print("figure:", p)
    deck = build_deck(figs)
    print("deck  :", deck)


if __name__ == "__main__":
    main()
